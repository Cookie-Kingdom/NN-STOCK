from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "docs" / "คู่มือรหัสเอกสารและล็อตสำหรับลูกค้า.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, color=None, align=None, size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(3)
    if align is not None:
        p.alignment = align
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Leelawadee UI"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Leelawadee UI")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    borders(cell)


def add_table(document, headers, data, widths):
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    for i, title in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Cm(widths[i])
        shade(cell, "1B5A4B")
        set_cell_text(cell, title, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
    for row_index, row in enumerate(data):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = Cm(widths[i])
            if row_index % 2:
                shade(cells[i], "F4F7F5")
            set_cell_text(cells[i], value, size=9)
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def heading(document, text, level=1):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Leelawadee UI"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Leelawadee UI")
    run.font.size = Pt(15 if level == 1 else 12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def paragraph(document, text, bold_prefix=None):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.22
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.name = "Leelawadee UI"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Leelawadee UI")
        r.font.size = Pt(10.5)
        text = text[len(bold_prefix):]
    r = p.add_run(text)
    r.font.name = "Leelawadee UI"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Leelawadee UI")
    r.font.size = Pt(10.5)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.7)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)

styles = doc.styles
styles["Normal"].font.name = "Leelawadee UI"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Leelawadee UI")
styles["Normal"].font.size = Pt(10.5)

header = section.header.paragraphs[0]
header.text = "NerdNuea Stock"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.runs[0].font.name = "Leelawadee UI"
header.runs[0]._element.rPr.rFonts.set(qn("w:eastAsia"), "Leelawadee UI")
header.runs[0].font.size = Pt(9)
header.runs[0].font.color.rgb = RGBColor(80, 104, 95)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("คู่มืออ้างอิงรหัสเอกสารและล็อต")
footer.runs[0].font.name = "Leelawadee UI"
footer.runs[0]._element.rPr.rFonts.set(qn("w:eastAsia"), "Leelawadee UI")
footer.runs[0].font.size = Pt(8.5)
footer.runs[0].font.color.rgb = RGBColor(100, 110, 105)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(5)
r = title.add_run("คู่มือรหัสเอกสารและล็อต")
r.font.name = "Leelawadee UI"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "Leelawadee UI")
r.font.size = Pt(24)
r.font.color.rgb = RGBColor(0, 0, 0)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(16)
run = subtitle.add_run("สำหรับลูกค้าและคู่ค้าที่ใช้งาน NerdNuea Stock")
run.font.name = "Leelawadee UI"
run._element.rPr.rFonts.set(qn("w:eastAsia"), "Leelawadee UI")
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(83, 107, 97)

paragraph(doc, "เอกสารนี้อธิบายความหมายของรหัสที่พบในใบสั่งซื้อ ใบขนส่ง ใบ Invoice และบันทึกการผลิต เพื่อให้ทุกฝ่ายอ้างอิงรายการเดียวกันได้ตลอดเส้นทางตั้งแต่ Foodiva ส่งเนื้อ โรงรมควันผลิต จนกลับเข้าสต๊อกกลาง")
paragraph(doc, "รหัสแต่ละชุดใช้สำหรับติดตามรายการในระบบ ไม่ใช่เลขทะเบียนบริษัท เลขผู้เสียภาษี หรือหลักฐานการชำระเงิน โดยเอกสารที่มีผลทางบัญชีให้ยึดเลข Invoice และไฟล์แนบของผู้ขายเป็นหลัก")

heading(doc, "สรุปรหัสที่ใช้ในระบบ")
add_table(doc,
    ["รหัส", "ชื่อเอกสารหรือรายการ", "ใครเป็นผู้สร้าง", "ใช้ทำอะไร", "ตัวอย่าง"],
    [
        ["PO", "ใบสั่งซื้อเนื้อ", "Owner", "สั่งซื้อเนื้อจาก Foodiva และเป็นจุดเริ่มต้นของรายการ", "PO-2026-0001"],
        ["F", "Lot เนื้อหลัก", "ระบบเมื่อ Owner ออก PO", "ระบุเนื้อชุดเดียวกันตลอดกระบวนการ", "F260913-001"],
        ["Invoice Foodiva", "Invoice เนื้อ", "Foodiva", "ยืนยันน้ำหนักและยอดเนื้อที่ส่งจริงตาม PO", "เลข Invoice ของ Foodiva"],
        ["SO", "ใบสั่ง PO โรงรมควัน", "Owner", "สั่งบริการรมควันให้ Chef_house โดยอ้างอิง Lot เนื้อ", "SO-2026-0001"],
        ["Invoice Chef_house", "Invoice ค่ารมควัน", "Chef_house", "แจ้งยอดค่ารมควันเพื่อให้ Owner ตรวจและชำระ", "เลข Invoice ของ Chef_house"],
        ["TR", "ใบขนส่ง", "Owner", "ติดตามรถ ต้นทาง ปลายทาง น้ำหนัก และคนขับ", "TR-2026-0001"],
        ["SB", "Lot สโมครายวัน", "ระบบเมื่อ Chef_house บันทึกการสโมค", "แยกรอบสโมคย่อยภายใต้ Lot เนื้อหลัก", "SB-2026-0001"],
    ],
    [2.3, 3.4, 2.7, 6.1, 3.0],
)

heading(doc, "วิธีอ่านรหัส")
paragraph(doc, "รหัส PO SO TR และ SB ใช้ปี ค ศ สี่หลัก ตามด้วยลำดับรายการสี่หลัก เช่น PO-2026-0001 หมายถึงใบสั่งซื้อเนื้อรายการที่ 1 ของปี 2026")
paragraph(doc, "รหัส Lot เนื้อหลักขึ้นต้นด้วย F ตามด้วยวันที่สร้างในรูปแบบ YYMMDD และลำดับสามหลัก เช่น F260913-001 หมายถึง Lot เนื้อที่สร้างวันที่ 13 กันยายน 2026 ลำดับที่ 001")
paragraph(doc, "รหัส Lot สโมคขึ้นต้นด้วย SB เป็นรอบการผลิตย่อย เช่น SB-2026-0002 คือรอบสโมคที่ 2 ซึ่งยังอ้างอิงกลับไปยัง Lot เนื้อหลักเดิมได้")

doc.add_page_break()
heading(doc, "ลำดับเอกสารตามการทำงาน")
add_table(doc,
    ["ลำดับ", "เอกสารหรือบันทึก", "การอ้างอิงที่ต้องตรวจ", "ผู้รับผิดชอบ"],
    [
        ["1", "PO เนื้อ", "เลข PO และ Lot เนื้อหลัก", "Owner"],
        ["2", "Invoice Foodiva", "เลข PO Lot น้ำหนักตาม Invoice และน้ำหนักพร้อมส่ง", "Foodiva"],
        ["3", "PO โรงรมควัน", "เลข SO และ Lot เนื้อหลัก อ้างอิง Invoice Foodiva", "Owner"],
        ["4", "Invoice Chef_house", "เลข SO Lot เนื้อ และยอดค่ารมควัน", "Chef_house และ Owner"],
        ["5", "ใบขนส่งขาไป", "เลข TR ต้นทาง Foodiva ปลายทาง Chef_house และน้ำหนักส่ง", "Owner"],
        ["6", "ใบรับเนื้อ Chef_house", "Lot เนื้อ น้ำหนักรับจริง และเวลารับ", "Chef_house"],
        ["7", "Lot สโมครายวัน", "เลข SB น้ำหนักเข้าเตา หลังรม Waste และจำนวนถุง", "Chef_house"],
        ["8", "ใบขนส่งขากลับ", "เลข TR ต้นทาง Chef_house ปลายทาง Foodiva และน้ำหนักหลังรม", "Owner"],
    ],
    [1.3, 4.2, 8.0, 4.0],
)

heading(doc, "หลักการอ้างอิงเมื่อมีคำถามหรือพบส่วนต่าง")
paragraph(doc, "ให้แจ้งเลข Lot เนื้อหลักก่อนเสมอ เพราะเป็นรหัสเดียวที่เชื่อม PO Invoice ใบขนส่ง และรอบสโมคทั้งหมดเข้าด้วยกัน ตัวอย่างเช่น หากต้องการตรวจน้ำหนักของ F260913-001 ให้เปิดเอกสารและ Traceability แล้วกดดาวน์โหลดเอกสารของ Lot นั้น")
paragraph(doc, "หากตรวจสอบเฉพาะรอบสโมค ให้แจ้งรหัส SB เพิ่มเติม เช่น SB-2026-0002 ระบบจะแสดงน้ำหนักเข้าเตา น้ำหนักหลังรม น้ำหนัก Waste และจำนวนถุงของรอบนั้น")
paragraph(doc, "หากเป็นเรื่องรถหรือการรับส่ง ให้แจ้งเลข TR พร้อม Lot เนื้อ เพื่อเทียบน้ำหนักส่งกับน้ำหนักรับและค้นหาทะเบียนรถหรือข้อมูลคนขับได้")

heading(doc, "ข้อควรทราบ")
paragraph(doc, "เลข Invoice ของ Foodiva และ Chef_house เป็นเลขที่คู่ค้าเป็นผู้ระบุ ระบบจัดเก็บเพื่อเชื่อมกับ PO และ Lot ไม่ได้ออกเลข Invoice แทนคู่ค้า")
paragraph(doc, "รหัสที่ขึ้นในหน้าจอเป็นข้อมูลอ้างอิงของระบบ การพิมพ์หรือดาวน์โหลด PDF จากหน้าเอกสารและ Traceability ใช้สำหรับตรวจสอบเส้นทางเอกสาร ส่วนไฟล์ Invoice ที่คู่ค้าแนบให้ดาวน์โหลดได้จากเมนูใบ Invoice")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
